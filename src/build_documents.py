from pathlib import Path
import csv,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
P=Path(__file__).resolve().parents[1]/'publication'

def newdoc():
 d=Document(); s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.85)
 for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
  st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
  st.font.size=Pt(11 if name=='Normal' else 18 if name=='Title' else 13 if name=='Heading 1' else 11)
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
 for st in d.styles:
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
  for fonts in st.element.iter(qn('w:rFonts')):
   for attr in list(fonts.attrib):
    if 'Theme' in attr:del fonts.attrib[attr]
 d.styles['Title'].paragraph_format.space_after=Pt(12)
 for name in ['Heading 1','Heading 2']:d.styles[name].paragraph_format.space_before=Pt(10)
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 return d

def table(d,title,heads,rows,widths):
 p=d.add_paragraph(title,style='Caption');p.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1, cols=len(heads));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for c,w,h in zip(t.rows[0].cells,widths,heads):c.width=Inches(w);c.text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,v,w in zip(cells,row,widths):c.text=str(v);c.width=Inches(w)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for column_index,c in enumerate(row.cells):
   cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   cp.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');cp.append(sh)
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if column_index>0 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.size=Pt(9);r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(1)

def records(n):return list(csv.DictReader((P/'tables'/n).open()))
def expand(d,token):
 if token=='[TABLE1]':
  rows=[[r['seed'],f"{float(r['validation_gap']):.2f}",f"{float(r['test_gap']):.2f}",f"{float(r['profit']):,.0f}",f"{float(r['difference']):+.0f}"] for r in records('split_summary.csv')]
  table(d,'Table 1 Validation and test absolute senior TPR gaps and constrained policy profit',['Seed','Validation gap (pp)','Test gap (pp)','EO profit (CU)','EO minus global (CU)'],rows,[.6,1.4,1.2,1.4,2.2])
 elif token=='[TABLE2]':
  names={'catboost':'CatBoost','random_forest':'Random forest','logistic':'Logistic'}
  rows=[]
  for r in records('model_summary.csv'):
   name=r['model'];family=next(k for k in names if name.startswith(k));label=name[len(family)+1:].replace('_original','').replace('_',' ')
   rows.append([names[family]+' '+label,f"{float(r['roc_auc']):.4f}",f"{float(r['brier']):.4f}",f"{float(r['scenario_profit']):,.1f}"])
  table(d,'Table 2 Mean calibrated test metrics across ten overlapping splits at the assumed Bayes cutoff',['Configuration','ROC AUC','Brier score','Scenario profit (CU)'],rows,[3.2,1,1,1.6])
 elif token=='[TABLE3]':
  rows=[[f"{float(r['clv']):.0f}",f"{float(r['acceptance']):.2f}",r['failures']+'/10',f"{float(r['median_gap_pp']):.2f}"] for r in records('economics_summary.csv')]
  table(d,'Table 3 Economic sensitivity of the constrained policy on test data',['Customer value (CU)','Acceptance','Gap exceedances','Median absolute gap (pp)'],rows,[1.6,1.2,1.7,2.3])
 else:
  n=1 if token=='[FIGURE1]' else 2
  p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(P/f'figures/figure{n}.png'),width=Inches(6.6))
  cap='Fig. 1 Absolute senior-group TPR gaps on validation and test partitions. The dotted line marks the five-percentage-point tolerance. Repetitions overlap and are not independent datasets.' if n==1 else 'Fig. 2 Constrained-minus-global test scenario profit differences with conditional 95% paired bootstrap intervals. CU denotes assumed currency units; models, policies, and demographic/outcome counts remain fixed within each interval.'
  d.add_paragraph(cap,style='Caption')

d=newdoc()
for chunk in (P/'Manuscript.md').read_text().split('\n\n'):
 chunk=chunk.strip()
 if not chunk:continue
 if chunk.startswith('[TABLE') or chunk.startswith('[FIGURE'):expand(d,chunk)
 elif chunk.startswith('# '):d.add_paragraph(chunk[2:],'Title')
 elif chunk.startswith('### '):d.add_paragraph(chunk[4:],'Heading 2')
 elif chunk.startswith('## '):d.add_paragraph(chunk[3:],'Heading 1')
 else:d.add_paragraph(chunk)
d.save(P/'Manuscript.docx')
d=newdoc();d.add_paragraph('Cover letter for SN Computer Science','Title')
for text in ['[Submission date]','Dear Editors,','Please consider our research article, “Validation to test transfer of equal opportunity constraints in profit sensitive telecom churn prediction,” for publication in SN Computer Science.',
'The study evaluates whether a profit-maximizing senior-group recall constraint selected on validation data transfers to held-out data in the public IBM Telco Customer Churn benchmark. Fitting, calibration, policy selection and testing use separate partitions. Across ten recorded repetitions, all constrained policies met the validation tolerance, while seven exceeded it on test data.',
'The contribution is a reproducible empirical evaluation of constraint transfer and economic sensitivity. It does not claim a new algorithm, observed campaign savings, causal retention uplift, or external validation. We explicitly acknowledge related profit-and-fairness research and the retrospective nature of the ten-split extension. The accompanying materials contain code, measured results and provenance information.',
'The manuscript addresses applied machine-learning evaluation and responsible decision support. We request the subscription publication route, subject to applicable author and funder requirements.',
'[Before submission, insert verified statements about prior publication and concurrent review, approval by all authors, and any related manuscripts or preprints. These facts have not yet been supplied.]',
'[Complete funding, competing interests, author contributions, ethics requirements and corresponding-author details in the manuscript. Review and finalize the AI-assistance disclosure.]',
'Sincerely,','[Corresponding author full name]','[Affiliation]','[Email]']:
 d.add_paragraph(text)
d.save(P/'Cover_Letter.docx')
print('Created two Word files')

import hashlib,json
inputs=[P/'Manuscript.md',Path(__file__),*sorted((P/'tables').glob('*.csv')),*sorted((P/'figures').glob('*.png'))]
record={'inputs':{str(f.relative_to(P.parent)):hashlib.sha256(f.read_bytes()).hexdigest() for f in inputs},'outputs':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in [P/'Manuscript.docx',P/'Cover_Letter.docx']}}
(P/'document_build.json').write_text(json.dumps(record,indent=2))
